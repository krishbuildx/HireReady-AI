"""Resume text extraction for PDF, DOCX, TXT."""
from __future__ import annotations

import io
import re
from typing import Dict, List

import fitz  # PyMuPDF
import pdfplumber
from docx import Document


def _extract_pdf(data: bytes) -> str:
    text = ""
    # Primary: PyMuPDF
    try:
        with fitz.open(stream=data, filetype="pdf") as doc:
            text = "\n".join(page.get_text("text") for page in doc)
    except Exception:
        text = ""
    # Fallback: pdfplumber
    if not text.strip():
        try:
            with pdfplumber.open(io.BytesIO(data)) as pdf:
                text = "\n".join((p.extract_text() or "") for p in pdf.pages)
        except Exception:
            pass
    return text


def _extract_docx(data: bytes) -> str:
    doc = Document(io.BytesIO(data))
    parts: List[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def _extract_txt(data: bytes) -> str:
    for enc in ("utf-8", "latin-1"):
        try:
            return data.decode(enc)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="ignore")


def extract_text(filename: str, data: bytes) -> str:
    """Return extracted text for a resume file."""
    name = filename.lower()
    if name.endswith(".pdf"):
        text = _extract_pdf(data)
    elif name.endswith(".docx"):
        text = _extract_docx(data)
    elif name.endswith(".txt"):
        text = _extract_txt(data)
    else:
        raise ValueError(f"Unsupported file type: {filename}")
    # Normalize whitespace
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


SECTION_PATTERNS = {
    "summary": r"(professional\s+summary|summary|objective|profile)",
    "experience": r"(experience|employment|work\s+history|professional\s+experience)",
    "education": r"(education|academic)",
    "skills": r"(skills|technical\s+skills|core\s+competencies)",
    "projects": r"(projects|personal\s+projects|selected\s+projects)",
    "certifications": r"(certifications|licenses)",
    "achievements": r"(achievements|awards|honors)",
    "contact": r"(contact|email|phone)",
}


def detect_sections(text: str) -> Dict[str, str]:
    """Rule-based section splitter. Returns {section: content}."""
    lines = text.splitlines()
    section_map: Dict[str, List[str]] = {}
    current = "header"
    section_map[current] = []
    header_regex = re.compile(
        r"^\s*(" + "|".join(p for p in SECTION_PATTERNS.values()) + r")\s*[:\-]?\s*$",
        re.IGNORECASE,
    )
    for line in lines:
        stripped = line.strip()
        matched = None
        if stripped and len(stripped) < 60 and header_regex.match(stripped):
            for key, pat in SECTION_PATTERNS.items():
                if re.search(pat, stripped, re.IGNORECASE):
                    matched = key
                    break
        if matched:
            current = matched
            section_map.setdefault(current, [])
        else:
            section_map.setdefault(current, []).append(line)
    return {k: "\n".join(v).strip() for k, v in section_map.items() if "\n".join(v).strip()}


def missing_sections(sections: Dict[str, str]) -> List[str]:
    required = ["summary", "experience", "education", "skills"]
    return [s for s in required if s not in sections]
